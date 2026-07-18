from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from html import unescape
from io import BytesIO
from typing import Any

import psycopg
import requests
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_core.messages import HumanMessage, SystemMessage
from pydantic import BaseModel, create_model
from psycopg.rows import dict_row
from psycopg.types.json import Json
from pypdf import PdfReader
from rapidfuzz import fuzz

try:
    # Package-submodule context (backend's INGEST_LOCAL_INVOKE local-invoke path):
    # this file loads as catalyst_functions.ingest_processor.pipeline.processor, so
    # ..llm correctly resolves to the sibling ingest_processor/llm.py.
    from ..llm import build_classifier, build_extractor, describe_image
except ImportError:
    # Deployed Catalyst runtime: main.py sits flat at /catalyst/ with pipeline/ as a
    # sibling dir and no parent package, so `pipeline` itself loads as a top-level
    # package -- `..llm` tries to go above that (ImportError: attempted relative
    # import beyond top-level package). llm.py is a flat sibling of pipeline/ there,
    # so the plain top-level name resolves via sys.path instead (confirmed live,
    # same pattern as main.py's own .pipeline.processor fallback).
    from llm import build_classifier, build_extractor, describe_image


logger = logging.getLogger(__name__)


# Log all expected env vars at import time so Catalyst portal logs immediately show
# which vars are missing, without having to trigger a full run first.
def _audit_env() -> None:
    required = [
        "DB_HOST", "DB_PORT", "DB_USER", "DB_PASSWORD", "DB_NAME",
        "ZOHO_CATALYST_CLIENT_ID", "ZOHO_CATALYST_CLIENT_SECRET", "ZOHO_CATALYST_REFRESH_TOKEN",
        "ZOHO_CATALYST_PROJECT_ID", "ZOHO_CATALYST_PROJECT_KEY",
        "ZOHO_CATALYST_AUTH_DOMAIN", "ZOHO_CATALYST_PROJECT_DOMAIN",
        "ZOHO_STRATUS_BUCKET",
        "ZOHO_QUICKML_ENDPOINT_URL", "ZOHO_QUICKML_CATALYST_ORG", "ZOHO_QUICKML_MODEL_NAME",
    ]
    optional = [
        "NEO4J_URI", "NEO4J_USERNAME", "NEO4J_PASSWORD",
        "PINECONE_API_KEY", "PINECONE_INDEX",
        "AWS_DEFAULT_REGION", "BEDROCK_EMBEDDING_MODEL",
        "SPLINK_ENDPOINT_URL", "DATA_INGESTION_LLM",
        "ZOHO_QUICKML_VLM_ENDPOINT_URL",
    ]
    missing = [k for k in required if not os.getenv(k, "").strip()]
    present_optional = [k for k in optional if os.getenv(k, "").strip()]
    if missing:
        logger.error("env_audit: MISSING required vars: %s", missing)
    else:
        logger.info("env_audit: all %d required vars present", len(required))
    logger.info("env_audit: optional vars set: %s", present_optional)


_audit_env()


_ACCOUNT_RE = re.compile(r"\b\d{9,20}\b")
_PHONE_RE = re.compile(r"\b(?:\+91[-\s]?)?[6-9]\d{9}\b")
_UPI_RE = re.compile(r"\b[a-zA-Z0-9.\-_]{2,}@[a-zA-Z]{2,}\b")
_IMEI_RE = re.compile(r"\b\d{14,16}\b")
_LABEL_SANITIZE_RE = re.compile(r"[^A-Za-z0-9_]")

# pole_entity_type == 'Object' groups always resolve their single hard identifier
# through one of these four tables; the non-identifier fields of the same group
# update other columns on the same row (see _write_object_row).
OBJECT_TABLE = {
    "Account": {"raw": "account_number_raw", "norm": "account_number_normalized", "pk": "account_id", "identifier_type": "account_number"},
    "UPIHandle": {"raw": "vpa_raw", "norm": "vpa_normalized", "pk": "upi_id", "identifier_type": "upi"},
    "PhoneNumber": {"raw": "number_raw", "norm": "number_normalized", "pk": "phone_id", "identifier_type": "phone"},
    "Device": {"raw": "imei_raw", "norm": "imei_normalized", "pk": "device_id", "identifier_type": "imei"},
}

# pole_entity_type == 'Person' groups: KSP's PKs are plain INTEGER (not autoincrement),
# so we mint the next id manually, same convention CaseMaster already used.
PERSON_TABLE = {
    "Accused": {"pk": "AccusedMasterID", "name_column": "AccusedName"},
    "Victim": {"pk": "VictimMasterID", "name_column": "VictimName"},
    "ComplainantDetails": {"pk": "ComplainantID", "name_column": "ComplainantName"},
}


def _env(key: str, default: str = "") -> str:
    return os.getenv(key, default).strip()


def _norm(identifier_type: str, value: str) -> str:
    value = (value or "").strip()
    if identifier_type == "account_number":
        return "".join(value.split())
    if identifier_type == "phone":
        digits = "".join(ch for ch in value if ch.isdigit())
        return digits[-10:] if len(digits) >= 10 else digits
    if identifier_type == "upi":
        return value.lower()
    if identifier_type == "imei":
        return "".join(ch for ch in value if ch.isdigit())
    return value


def _utc_now() -> str:
    return datetime.now(tz=timezone.utc).isoformat()


def _normalize_timestamp(raw: Any) -> str:
    """The LLM sometimes extracts a bare time (e.g. a chat screenshot's "10:05")
    instead of a full timestamp -- postgres rejects that for a timestamptz column.
    Fall back to now() rather than fail the whole file's load over one field."""
    if not raw:
        return _utc_now()
    text_value = str(raw).strip()
    try:
        datetime.fromisoformat(text_value.replace("Z", "+00:00"))
        return text_value
    except ValueError:
        logger.warning("normalize_timestamp: unparseable value %r — using now()", raw)
        return _utc_now()


def _sanitize_label(label: str) -> str:
    return _LABEL_SANITIZE_RE.sub("", label) or "Object"


def _build_database_url() -> str:
    if _env("DATABASE_URL"):
        logger.debug("db_url: using DATABASE_URL env var")
        return _env("DATABASE_URL")
    host = _env("DB_HOST")
    port = _env("DB_PORT", "5432")
    user = _env("DB_USER")
    name = _env("DB_NAME", "ksp_crime")
    ssl = _env("DB_SSL", "require")
    if not host:
        logger.error("db_url: DB_HOST is not set — connection will fail")
    if not user:
        logger.error("db_url: DB_USER is not set — connection will fail")
    logger.debug("db_url: host=%s port=%s user=%s dbname=%s sslmode=%s", host, port, user, name, ssl)
    password = _env("DB_PASSWORD").strip("\"'")
    # prepare_threshold is a psycopg.connect() kwarg, not a libpq/psycopg URI
    # parameter -- embedding it in the query string raises "invalid URI query
    # parameter" at connect time. Callers pass prepare_threshold=None explicitly.
    return f"postgresql://{user}:{password}@{host}:{port}/{name}?sslmode={ssl}"


def _patch_zcatalyst_url_join() -> None:
    from zcatalyst_sdk._http_client import HttpClient

    if getattr(HttpClient.request, "_ksp_patched", False):
        return
    original = HttpClient.request

    def _request(self, method, url=None, path=None, *args, **kwargs):
        if url is None and path and str(path).startswith("/oauth"):
            path = path.lstrip("/")
        return original(self, method, url, path, *args, **kwargs)

    _request._ksp_patched = True  # type: ignore[attr-defined]
    HttpClient.request = _request  # type: ignore[method-assign]
    logger.debug("zcatalyst_sdk: HttpClient OAuth path patch applied")


def _init_catalyst_app():
    logger.debug("catalyst_init: starting")
    for key, default in [
        ("X_ZOHO_CATALYST_ACCOUNTS_URL", _env("ZOHO_CATALYST_AUTH_DOMAIN", "https://accounts.zohoportal.in")),
        ("X_ZOHO_CATALYST_CONSOLE_URL", _env("ZOHO_CATALYST_PROJECT_DOMAIN", "https://api.catalyst.zoho.in")),
        ("X_ZOHO_STRATUS_RESOURCE_SUFFIX", ".zohostratus.in"),
    ]:
        os.environ.setdefault(key, default)
        logger.debug("catalyst_init: %s=%s", key, os.environ[key])

    import zcatalyst_sdk
    from zcatalyst_sdk import credentials, types

    _patch_zcatalyst_url_join()
    try:
        app = zcatalyst_sdk.get_app()
        logger.debug("catalyst_init: reused existing app")
        return app
    except Exception as exc:
        logger.debug("catalyst_init: get_app() failed (%s) — initializing fresh", exc)

    project_id = _env("ZOHO_CATALYST_PROJECT_ID")
    project_key = _env("ZOHO_CATALYST_PROJECT_KEY")
    environment = _env("ZOHO_CATALYST_ENVIRONMENT", "Development")
    project_domain = _env("ZOHO_CATALYST_PROJECT_DOMAIN", "https://api.catalyst.zoho.in")
    if not project_id or not project_key:
        logger.error("catalyst_init: ZOHO_CATALYST_PROJECT_ID or ZOHO_CATALYST_PROJECT_KEY not set")
    logger.debug("catalyst_init: project_id=%s project_key=%s environment=%s domain=%s",
                 project_id, project_key, environment, project_domain)

    cred = credentials.RefreshTokenCredential(
        {
            "refresh_token": _env("ZOHO_CATALYST_REFRESH_TOKEN"),
            "client_id": _env("ZOHO_CATALYST_CLIENT_ID"),
            "client_secret": _env("ZOHO_CATALYST_CLIENT_SECRET"),
        }
    )
    options = types.ICatalystOptions(
        project_id=project_id,
        project_key=project_key,
        environment=environment,
        project_domain=project_domain,
    )
    try:
        app = zcatalyst_sdk.initialize_app(credential=cred, options=options)
        logger.debug("catalyst_init: initialized app ok")
        return app
    except Exception as exc:
        logger.error("catalyst_init: initialize_app failed: %s", exc)
        raise


def _stratus_bucket():
    bucket_name = _env("ZOHO_STRATUS_BUCKET", "ksp-data-files")
    logger.debug("stratus_bucket: bucket=%s", bucket_name)
    try:
        app = _init_catalyst_app()
        bucket = app.stratus().bucket(bucket_name)
        logger.debug("stratus_bucket: ok domain=%s", getattr(bucket, "bucket_domain", "<unknown>"))
        return bucket
    except Exception as exc:
        logger.error("stratus_bucket: failed to get bucket=%s: %s", bucket_name, exc)
        raise


@dataclass
class RunParams:
    batch_id: str
    case_id: int
    run_id: str
    phase: str  # "extract" (Phase A) | "load" (Phase B)


class IngestProcessor:
    def __init__(self, params: RunParams):
        self.params = params
        logger.info(
            "IngestProcessor init run_id=%s batch_id=%s case_id=%s phase=%s  db_host=%s db_port=%s db_user=%s",
            params.run_id, params.batch_id, params.case_id, params.phase,
            _env("DB_HOST", "<not set>"), _env("DB_PORT", "5432"), _env("DB_USER", "<not set>"),
        )
        db_url = _build_database_url()
        try:
            self.pg = psycopg.connect(db_url, row_factory=dict_row, prepare_threshold=None)
            logger.debug("db_connect: ok server_version=%s", self.pg.info.server_version)
        except Exception as exc:
            logger.error("db_connect: FAILED host=%s port=%s user=%s dbname=%s — %s",
                         _env("DB_HOST"), _env("DB_PORT", "5432"), _env("DB_USER"), _env("DB_NAME", "ksp_crime"), exc)
            raise
        self._lookup_cache: dict[str, dict[str, int]] = {}

    def close(self) -> None:
        self.pg.close()

    # ------------------------------------------------------------------
    # Master-data lookups (Occupation/Caste/Religion are INTEGER FKs, but the LLM
    # extractor returns free text -- e.g. "Retired Professor" for OccupationID --
    # so free text must be resolved against the master table before INSERT or
    # psycopg raises InvalidTextRepresentation. Columns are nullable FKs, so an
    # unresolvable value degrades to NULL (logged) rather than failing the load.
    # ------------------------------------------------------------------

    _LOOKUP_TABLES: dict[str, tuple[str, str, str]] = {
        "OccupationID": ("OccupationMaster", "OccupationID", "OccupationName"),
        "CasteID": ("CasteMaster", "caste_master_id", "caste_master_name"),
        "ReligionID": ("ReligionMaster", "ReligionID", "ReligionName"),
    }

    def _resolve_lookup_id(self, column: str, raw_value: str) -> int | None:
        cache = self._lookup_cache.setdefault(column, {})
        if not cache:
            table, id_col, name_col = self._LOOKUP_TABLES[column]
            with self.pg.cursor() as cur:
                cur.execute(f"SELECT {id_col} AS id, {name_col} AS name FROM {table}")
                for row in cur.fetchall():
                    cache[str(row["name"]).strip().lower()] = int(row["id"])

        key = raw_value.strip().lower()
        if key in cache:
            return cache[key]

        best_name, best_score = None, 0.0
        for name in cache:
            score = fuzz.ratio(key, name)
            if score > best_score:
                best_name, best_score = name, score
        if best_name and best_score >= 85:
            logger.debug("lookup: fuzzy-matched %s=%r -> %r (score=%.0f)", column, raw_value, best_name, best_score)
            return cache[best_name]

        logger.warning("lookup: no match for %s=%r — leaving NULL", column, raw_value)
        return None

    # ------------------------------------------------------------------
    # Status / Signals
    # ------------------------------------------------------------------

    def _update_stage(self, stage: str, status: str | None = None, file: str | None = None, error_message: str | None = None) -> dict[str, Any]:
        with self.pg.cursor() as cur:
            if file:
                cur.execute(
                    """
                    UPDATE PipelineRun
                    SET current_stage = %s,
                        status = COALESCE(%s, status),
                        files_progress = jsonb_set(COALESCE(files_progress, '{}'::jsonb), %s, %s::jsonb, true),
                        error_message = %s,
                        updated_at = NOW()
                    WHERE run_id = %s
                    RETURNING files_progress
                    """,
                    (stage, status, [file], Json({"stage": stage, "status": status or "in_progress", "ts": _utc_now()}), error_message, self.params.run_id),
                )
            else:
                cur.execute(
                    """
                    UPDATE PipelineRun
                    SET current_stage = %s, status = COALESCE(%s, status), error_message = %s, updated_at = NOW()
                    WHERE run_id = %s
                    RETURNING files_progress
                    """,
                    (stage, status, error_message, self.params.run_id),
                )
            row = cur.fetchone()
        self.pg.commit()
        files_progress = (row or {}).get("files_progress") or {}
        self._publish_signal(stage=stage, status=status, file=file, files_progress=files_progress)
        return files_progress

    def _publish_signal(self, stage: str, status: str | None, file: str | None, files_progress: dict[str, Any]) -> None:
        url = _env("SIGNALS_PIPELINE_PUBLISHER_URL")
        if not url:
            logger.debug("publish_signal: SIGNALS_PIPELINE_PUBLISHER_URL not set — skipping signal stage=%s", stage)
            return
        try:
            resp = requests.post(
                url,
                json={
                    "run_id": self.params.run_id,
                    "stage": stage,
                    "status": status,
                    "file": file,
                    "files_progress": files_progress,
                    "ts": _utc_now(),
                },
                timeout=2,
            )
            # This is a best-effort push (the WS's Postgres poll fallback covers a
            # missed signal), but a non-2xx response was previously invisible --
            # requests.post() doesn't raise on its own, so a wrong digest/expired
            # auth on the Signals publisher URL would silently degrade every run
            # to poll-only with no log line anywhere pointing at why.
            if not resp.ok:
                logger.warning(
                    "signals publish non-2xx run_id=%s stage=%s status_code=%s body=%s",
                    self.params.run_id, stage, resp.status_code, resp.text[:300],
                )
        except Exception as exc:
            logger.warning("signals publish failed run_id=%s stage=%s: %s", self.params.run_id, stage, exc)

    # ------------------------------------------------------------------
    # Stratus (raw / processed / archive)
    # ------------------------------------------------------------------

    def _list_raw_files(self) -> list[dict[str, str]]:
        bucket = _stratus_bucket()
        prefix = f"raw/{self.params.batch_id}/"
        logger.debug("list_raw_files: prefix=%s", prefix)
        files: list[dict[str, str]] = []
        next_token = None
        page_num = 0
        while True:
            page_num += 1
            try:
                page = bucket.list_paged_objects(prefix=prefix, next_token=next_token)
            except Exception as exc:
                logger.error("list_raw_files: stratus list failed prefix=%s page=%d: %s", prefix, page_num, exc)
                raise
            for obj in page["contents"]:
                key = obj.to_dict().get("key") if hasattr(obj, "to_dict") else obj.get("key")
                if not key:
                    continue
                remainder = key[len(prefix):]
                file_type, _, filename = remainder.partition("_")
                files.append({"key": key, "file_type": file_type or "evidence", "filename": filename or remainder})
            if not page.get("truncated"):
                break
            next_token = page.get("next_continuation_token")
        logger.info("list_raw_files: found %d file(s) under %s", len(files), prefix)
        return files

    def _get_bytes(self, key: str) -> bytes:
        logger.debug("get_bytes: key=%s", key)
        bucket = _stratus_bucket()
        try:
            obj = bucket.get_object(key)
        except Exception as exc:
            logger.error("get_bytes: stratus get_object failed key=%s: %s", key, exc)
            raise
        if isinstance(obj, bytes):
            logger.debug("get_bytes: ok key=%s size=%d", key, len(obj))
            return obj
        if hasattr(obj, "read"):
            data = obj.read()
            logger.debug("get_bytes: ok (stream) key=%s size=%d", key, len(data))
            return data
        raise RuntimeError("Unsupported Stratus object payload type.")

    def _checkpoint_prefix(self) -> str:
        return f"processed/{self.params.case_id}/{self.params.run_id}/"

    def _write_checkpoint(self, manifest: dict[str, Any]) -> None:
        bucket = _stratus_bucket()
        key = self._checkpoint_prefix() + "manifest.json"
        data = json.dumps(manifest, default=str).encode("utf-8")
        logger.debug("write_checkpoint: key=%s size=%d", key, len(data))
        try:
            bucket.put_object(key, data, options={"overwrite": "true"})
            logger.info("write_checkpoint: ok key=%s", key)
        except Exception as exc:
            logger.error("write_checkpoint: FAILED key=%s: %s", key, exc)
            raise

    def _read_checkpoint(self) -> dict[str, Any]:
        bucket = _stratus_bucket()
        key = self._checkpoint_prefix() + "manifest.json"
        logger.debug("read_checkpoint: key=%s", key)
        try:
            data = bucket.get_object(key)
        except Exception as exc:
            logger.error("read_checkpoint: FAILED key=%s: %s", key, exc)
            raise
        if isinstance(data, bytes):
            parsed = json.loads(data.decode("utf-8"))
        else:
            parsed = json.loads(data.read().decode("utf-8"))
        logger.debug("read_checkpoint: ok key=%s files=%d", key, len(parsed.get("files", [])))
        return parsed

    def _archive_raw_files(self) -> list[str]:
        bucket = _stratus_bucket()
        moved: list[str] = []
        for entry in self._list_raw_files():
            raw_key = entry["key"]
            archive_key = f"archive/{self.params.case_id}/{self.params.run_id}/" + raw_key.rsplit("/", 1)[-1]
            logger.debug("archive: %s -> %s", raw_key, archive_key)
            try:
                bucket.rename_object(raw_key, archive_key)
            except Exception as exc:
                logger.warning("archive: rename failed (%s) — falling back to copy+delete", exc)
                try:
                    bucket.copy_object(raw_key, archive_key)
                    bucket.delete_object(raw_key)
                except Exception as exc2:
                    logger.error("archive: copy+delete also failed for key=%s: %s", raw_key, exc2)
                    raise
            moved.append(archive_key)
        logger.info("archive: moved %d file(s)", len(moved))
        return moved

    # ------------------------------------------------------------------
    # Text extraction / classification
    # ------------------------------------------------------------------

    def _extract_text(self, key: str, payload: bytes) -> str:
        lower = key.lower()
        logger.debug("extract_text: key=%s size=%d", key, len(payload))
        if lower.endswith(".txt"):
            logger.debug("extract_text: format=txt")
            return payload.decode("utf-8", errors="ignore")
        if lower.endswith(".html") or lower.endswith(".htm"):
            logger.debug("extract_text: format=html")
            html = payload.decode("utf-8", errors="ignore")
            soup = BeautifulSoup(html, "html.parser")
            return unescape(soup.get_text("\n"))
        if lower.endswith(".pdf"):
            logger.debug("extract_text: format=pdf")
            reader = PdfReader(BytesIO(payload))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if lower.endswith(".docx"):
            logger.debug("extract_text: format=docx")
            document = DocxDocument(BytesIO(payload))
            return "\n".join(p.text for p in document.paragraphs if p.text)
        if lower.endswith((".png", ".jpg", ".jpeg", ".webp")):
            logger.debug("extract_text: format=image — calling VLM")
            return describe_image(payload)
        logger.debug("extract_text: format=unknown — treating as utf-8 text")
        return payload.decode("utf-8", errors="ignore")

    def _active_evidence_doc_types(self) -> list[dict[str, Any]]:
        with self.pg.cursor() as cur:
            cur.execute(
                "SELECT doc_type, description FROM SchemaDefinition WHERE is_active = true AND doc_type LIKE 'EVIDENCE_%' ORDER BY doc_type"
            )
            return cur.fetchall()

    def _classify(self, raw_text: str, file_type: str) -> str:
        if file_type == "fir":
            logger.debug("classify: file_type=fir — returning FIR directly")
            return "FIR"
        if file_type == "ir":
            logger.debug("classify: file_type=ir — returning IR directly")
            return "IR"

        candidates = self._active_evidence_doc_types()
        labels = [c["doc_type"] for c in candidates] or ["EVIDENCE_BANK_STATEMENT", "EVIDENCE_UPI_SCREENSHOT", "EVIDENCE_CHAT_SCREENSHOT"]
        prompt_lines = [f"- {c['doc_type']}: {c.get('description') or ''}" for c in candidates]
        logger.debug("classify: evidence file — LLM classification among %s", labels)
        classifier = build_classifier()
        try:
            response = classifier.invoke(
                [
                    SystemMessage(
                        content=(
                            "Classify this evidence document into exactly one of these labels:\n"
                            + "\n".join(prompt_lines or labels)
                        )
                    ),
                    HumanMessage(content=raw_text[:4000]),
                ]
            )
        except Exception as exc:
            logger.error("classify: LLM classifier failed — defaulting to %s: %s", labels[0], exc)
            return labels[0]
        content = str(getattr(response, "content", "")).upper()
        for candidate in labels:
            if candidate in content:
                logger.debug("classify: LLM returned doc_type=%s", candidate)
                return candidate
        logger.warning("classify: LLM response did not match any label (response=%s) — defaulting to %s", content[:100], labels[0])
        return labels[0]

    # ------------------------------------------------------------------
    # Schema loading + dynamic structured extraction (repeating groups -> nested lists)
    # ------------------------------------------------------------------

    def _load_schema(self, doc_type: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        logger.debug("load_schema: doc_type=%s", doc_type)
        with self.pg.cursor() as cur:
            cur.execute(
                """
                SELECT schema_id, doc_type, version, description
                FROM SchemaDefinition
                WHERE doc_type = %s AND is_active = true
                ORDER BY version DESC
                LIMIT 1
                """,
                (doc_type,),
            )
            schema = cur.fetchone()
            if not schema:
                logger.error("load_schema: no active schema found for doc_type=%s", doc_type)
                raise RuntimeError(f"No active schema for doc_type={doc_type}")

            cur.execute(
                """
                SELECT field_id, group_name, is_repeating_group, pole_entity_type, field_name, data_type,
                       is_required, target_table, target_column, is_identifier, identifier_type
                FROM SchemaField
                WHERE schema_id = %s
                ORDER BY group_name, COALESCE(display_order, 0), field_id
                """,
                (schema["schema_id"],),
            )
            fields = cur.fetchall()
        logger.debug("load_schema: found schema_id=%s version=%s fields=%d", schema["schema_id"], schema.get("version"), len(fields))
        return schema, fields

    @staticmethod
    def _fields_by_group(fields: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
        grouped: dict[str, list[dict[str, Any]]] = {}
        for field in fields:
            grouped.setdefault(field["group_name"], []).append(field)
        return grouped

    def _build_dynamic_model(self, fields: list[dict[str, Any]]) -> type[BaseModel]:
        py_type_map: dict[str, Any] = {"string": str, "integer": int, "float": float, "date": str, "boolean": bool}
        model_fields: dict[str, tuple[Any, Any]] = {}
        for group_name, group_fields in self._fields_by_group(fields).items():
            is_repeating = bool(group_fields[0].get("is_repeating_group"))
            if is_repeating:
                sub_fields: dict[str, tuple[Any, Any]] = {}
                for field in group_fields:
                    py_type = py_type_map.get(str(field.get("data_type", "string")).lower(), str)
                    sub_fields[field["field_name"]] = (py_type | None, None)
                sub_model = create_model(f"{_sanitize_label(group_name)}Item", **sub_fields)  # type: ignore[arg-type]
                model_fields[group_name] = (list[sub_model] | None, None)
            else:
                for field in group_fields:
                    py_type = py_type_map.get(str(field.get("data_type", "string")).lower(), str)
                    model_fields[field["field_name"]] = (py_type | None, None)
        return create_model("DynamicExtractModel", **model_fields)  # type: ignore[arg-type]

    def _regex_fallback_extract(self, raw_text: str, fields: list[dict[str, Any]]) -> dict[str, Any]:
        extracted: dict[str, Any] = {}
        for group_name, group_fields in self._fields_by_group(fields).items():
            is_repeating = bool(group_fields[0].get("is_repeating_group"))
            row: dict[str, Any] = {}
            for field in group_fields:
                name = field["field_name"]
                identifier_type = field.get("identifier_type")
                if identifier_type == "account_number":
                    matches = _ACCOUNT_RE.findall(raw_text)
                    row[name] = matches[0] if matches else None
                elif identifier_type == "phone":
                    matches = _PHONE_RE.findall(raw_text)
                    row[name] = matches[0] if matches else None
                elif identifier_type == "upi":
                    matches = _UPI_RE.findall(raw_text)
                    row[name] = matches[0] if matches else None
                elif identifier_type == "imei":
                    matches = _IMEI_RE.findall(raw_text)
                    row[name] = matches[0] if matches else None
                elif name in {"brief_facts", "findings_narrative", "message_texts"}:
                    row[name] = raw_text[:4000]
                else:
                    row[name] = None
            if is_repeating:
                # ponytail: regex fallback only ever recovers a single instance of a
                # repeating group (e.g. one Accused, one Transaction row), never several.
                # Ceiling: undercounts multi-item documents when the LLM call fails.
                # Upgrade path: none needed here — this path only runs if the LLM is down.
                if any(v is not None for v in row.values()):
                    extracted[group_name] = [row]
            else:
                extracted.update(row)
        return extracted

    def _extract_structured(self, raw_text: str, fields: list[dict[str, Any]]) -> dict[str, Any]:
        model = self._build_dynamic_model(fields)
        try:
            extractor = build_extractor(model)
            response = extractor.invoke(
                [
                    SystemMessage(content="Extract the requested fields as JSON, matching the schema exactly."),
                    HumanMessage(content=raw_text[:12000]),
                ]
            )
            if isinstance(response, BaseModel):
                data = response.model_dump()
            elif isinstance(response, dict):
                data = response
            else:
                raise RuntimeError("Unstructured extractor response.")
            # Normalize sub-model items (Pydantic objects) to plain dicts.
            for key, value in list(data.items()):
                if isinstance(value, list):
                    data[key] = [item.model_dump() if isinstance(item, BaseModel) else item for item in value]
            logger.debug("structured extract ok fields=%d", len(data))
            return data
        except Exception as exc:
            logger.warning("structured extract failed (%s); using regex fallback", exc)
            return self._regex_fallback_extract(raw_text, fields)

    # ------------------------------------------------------------------
    # Phase A — extract, transform, checkpoint, entity match
    # ------------------------------------------------------------------

    def run_phase_a(self) -> dict[str, Any]:
        self._update_stage("EXTRACT_START", status="RUNNING")
        raw_files = self._list_raw_files()
        logger.info("phase_a run_id=%s raw_files_found=%d", self.params.run_id, len(raw_files))
        if not raw_files:
            self._update_stage("EXTRACT_DONE", status="FAILED", error_message="No raw files found for batch.")
            return {"run_id": self.params.run_id, "status": "FAILED"}

        file_entries: list[dict[str, Any]] = []
        review_items_created = 0
        # Candidates extracted from files already processed *in this same run* --
        # e.g. the FIR's Accused mention -- so a later file's mention of the same
        # person (the IR's Accused mention) gets compared against it too, not just
        # against historically-committed EntityMap rows. Without this, two brand
        # new mentions of the same person within one run never get a chance to be
        # linked (neither is in EntityMap yet), and _write_person_row would
        # otherwise write two separate SQL rows for one person.
        same_run_candidates: list[dict[str, Any]] = []

        for entry in raw_files:
            filename = entry["filename"]
            logger.debug("phase_a processing file=%s file_type=%s", filename, entry["file_type"])
            self._update_stage("FETCHING_FILE", file=filename)
            payload = self._get_bytes(entry["key"])

            self._update_stage("EXTRACTING_TEXT", file=filename)
            raw_text = self._extract_text(entry["key"], payload)

            self._update_stage("CLASSIFYING", file=filename)
            doc_type = self._classify(raw_text, entry["file_type"])
            logger.debug("classified file=%s doc_type=%s", filename, doc_type)

            self._update_stage("LOADING_SCHEMA", file=filename)
            schema, fields = self._load_schema(doc_type)

            self._update_stage("STRUCTURED_EXTRACT", file=filename)
            extracted = self._extract_structured(raw_text, fields)

            self._update_stage("TRANSFORM", file=filename)
            groups, provisional_uids = self._transform_to_groups(extracted, fields)

            self._update_stage("ENTITY_MATCH", file=filename)
            created = self._request_person_review(
                doc_type=doc_type, groups=groups, provisional_uids=provisional_uids, same_run_candidates=same_run_candidates
            )
            review_items_created += created

            file_entries.append(
                {
                    "key": entry["key"],
                    "filename": filename,
                    "file_type": entry["file_type"],
                    "doc_type": doc_type,
                    "schema_id": int(schema["schema_id"]),
                    "raw_text": raw_text[:40000],
                    "groups": groups,
                    "provisional_uids": provisional_uids,
                }
            )
            self._update_stage("CHECKPOINTED", status=None, file=filename)

        manifest = {
            "run_id": self.params.run_id,
            "case_id": self.params.case_id,
            "batch_id": self.params.batch_id,
            "created_at": _utc_now(),
            "files": file_entries,
        }
        self._update_stage("WRITING_CHECKPOINT")
        self._write_checkpoint(manifest)

        final_status = "REVIEW_PENDING"
        with self.pg.cursor() as cur:
            cur.execute(
                "UPDATE PipelineRun SET phase = 'REVIEW', status = %s, current_stage = 'REVIEW_PENDING', updated_at = NOW() WHERE run_id = %s",
                (final_status, self.params.run_id),
            )
        self.pg.commit()
        self._publish_signal(stage="REVIEW_PENDING", status=final_status, file=None, files_progress={})

        return {"run_id": self.params.run_id, "status": final_status, "review_items_created": review_items_created}

    def _transform_to_groups(self, extracted: dict[str, Any], fields: list[dict[str, Any]]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, str]]:
        """Group extracted fields by SchemaField.group_name, normalizing identifiers.
        Returns (groups, provisional_uids) where provisional_uids maps "group#index" -> uuid
        for every Person-type row (minted now, resolved for real in Phase B)."""
        groups: dict[str, list[dict[str, Any]]] = {}
        provisional_uids: dict[str, str] = {}

        for group_name, group_fields in self._fields_by_group(fields).items():
            is_repeating = bool(group_fields[0].get("is_repeating_group"))
            pole_type = group_fields[0].get("pole_entity_type")
            raw_rows = extracted.get(group_name) if is_repeating else [extracted]
            raw_rows = [r for r in (raw_rows or []) if isinstance(r, dict)]
            if not is_repeating and not raw_rows:
                raw_rows = [{}]

            columns = {f["field_name"]: f["target_column"] for f in group_fields}
            identifier_field = next((f["field_name"] for f in group_fields if f.get("is_identifier")), None)
            identifier_type = next((f.get("identifier_type") for f in group_fields if f.get("is_identifier")), None)

            produced_rows: list[dict[str, Any]] = []
            for row in raw_rows:
                normalized: dict[str, Any] = {}
                for field in group_fields:
                    value = row.get(field["field_name"])
                    if value is not None and field.get("is_identifier"):
                        value = _norm(str(field.get("identifier_type") or ""), str(value))
                    normalized[field["field_name"]] = value
                produced_rows.append(
                    {
                        "fields": normalized,
                        "target_table": group_fields[0]["target_table"],
                        "columns": columns,
                        "identifier_field": identifier_field,
                        "identifier_type": identifier_type,
                    }
                )

            if pole_type == "Person":
                for idx, row in enumerate(produced_rows):
                    provisional_uids[f"{group_name}#{idx}"] = str(uuid.uuid4())

            groups[group_name] = produced_rows

        return groups, provisional_uids

    def _request_person_review(
        self,
        doc_type: str,
        groups: dict[str, list[dict[str, Any]]],
        provisional_uids: dict[str, str],
        same_run_candidates: list[dict[str, Any]],
    ) -> int:
        endpoint = _env("SPLINK_ENDPOINT_URL")
        secret = _env("SPLINK_SHARED_SECRET")
        if not endpoint or not secret:
            logger.debug("request_person_review: SPLINK_ENDPOINT_URL or SPLINK_SHARED_SECRET not set — skipping entity resolution")
            return 0

        with self.pg.cursor() as cur:
            cur.execute(
                """
                SELECT e.entity_uid,
                       COALESCE(a.AccusedName, v.VictimName, c.ComplainantName) AS name
                FROM EntityMap e
                LEFT JOIN Accused a ON e.sql_table = 'Accused' AND e.sql_pk = a.AccusedMasterID::text
                LEFT JOIN Victim v ON e.sql_table = 'Victim' AND e.sql_pk = v.VictimMasterID::text
                LEFT JOIN ComplainantDetails c ON e.sql_table = 'ComplainantDetails' AND e.sql_pk = c.ComplainantID::text
                WHERE e.status = 'active' AND e.entity_type = 'Person'
                LIMIT 200
                """
            )
            existing = [dict(r) for r in cur.fetchall()]

        created_total = 0
        for group_name, group_meta in PERSON_TABLE.items():
            rows = groups.get(group_name) or []
            for idx, row in enumerate(rows):
                fields = row["fields"]
                name = _display_name(fields)
                if not name:
                    continue
                provisional_uid = provisional_uids.get(f"{group_name}#{idx}")
                candidate = {"entity_uid": provisional_uid, "group": group_name, "index": idx, "name": name}
                payload = {
                    "source_run_id": self.params.run_id,
                    "entity_type": "Person",
                    "candidate_record": candidate,
                    # existing (historical, already-committed) + same_run_candidates
                    # (mentions already seen from earlier files in *this* run, not
                    # yet in EntityMap) -- lets e.g. the IR's Accused mention match
                    # against the FIR's Accused mention from the same batch.
                    "existing_records": existing + same_run_candidates,
                    "persist_review_items": True,
                }
                try:
                    response = requests.post(
                        endpoint.rstrip("/") + "/internal/entity/splink-match",
                        # default=str: existing_records' entity_uid comes back from psycopg
                        # as a uuid.UUID, which requests' json= encoder can't serialize
                        # ("Object of type UUID is not JSON serializable"), silently failing
                        # every person match. Mirror the manifest write's json.dumps(default=str).
                        data=json.dumps(payload, default=str),
                        headers={"X-Splink-Secret": secret, "Content-Type": "application/json"},
                        timeout=30,
                    )
                    response.raise_for_status()
                    created_total += int(response.json().get("review_items_created", 0))
                except Exception as exc:
                    logger.warning("request_person_review: splink-match failed name=%s: %s", name, exc)
                    continue
                if provisional_uid:
                    same_run_candidates.append({"entity_uid": provisional_uid, "name": name})
        return created_total

    # ------------------------------------------------------------------
    # Phase B — apply review decisions, load SQL/Graph/Vector, archive
    # ------------------------------------------------------------------

    def run_phase_b(self) -> dict[str, Any]:
        self._update_stage("LOAD_START", status="RUNNING")
        logger.info("phase_b start run_id=%s", self.params.run_id)
        manifest = self._read_checkpoint()
        uid_remap = self._resolve_review_decisions()

        # FIR must land first so CaseMaster exists before any evidence/IR row's FK to it
        # (raw/ listing order is alphabetical by file_type prefix: evidence_ < fir_ < ir_).
        files_in_order = sorted(manifest["files"], key=lambda f: 0 if f["doc_type"] == "FIR" else 1)

        all_entity_uids: list[str] = []
        all_chunk_ids: list[str] = []

        for file_entry in files_in_order:
            filename = file_entry["filename"]
            logger.info("phase_b processing file=%s doc_type=%s", filename, file_entry.get("doc_type"))
            self._update_stage("SQL_LOAD", file=filename)
            groups, evidence_id = self._load_sql_groups(file_entry, uid_remap)

            self._update_stage("GRAPH_LOAD", file=filename)
            entity_uids = self._load_graph(file_entry["schema_id"], groups)
            all_entity_uids.extend(entity_uids)
            logger.debug("phase_b file=%s sql_load done, graph_load entity_uids=%d", filename, len(entity_uids))

            self._update_stage("VECTOR_LOAD", file=filename)
            chunk_ids = self._load_vector(file_entry, entity_uids)
            all_chunk_ids.extend(chunk_ids)
            logger.debug("phase_b file=%s vector_load chunk_ids=%d", filename, len(chunk_ids))

            self._update_stage("WRITEBACK", file=filename)
            self._writeback_links(chunk_ids, entity_uids)
            self._update_stage("FILE_DONE", file=filename)

        self._update_stage("ARCHIVING")
        archived = self._archive_raw_files()

        self._update_stage("DONE", status="COMPLETED")
        logger.info("phase_b done run_id=%s archived=%d", self.params.run_id, len(archived))
        with self.pg.cursor() as cur:
            cur.execute("UPDATE PipelineRun SET phase = 'DONE' WHERE run_id = %s", (self.params.run_id,))
        self.pg.commit()

        return {"run_id": self.params.run_id, "status": "COMPLETED", "archived_files": len(archived)}

    def _resolve_review_decisions(self) -> dict[str, str]:
        """Map provisional_entity_uid -> final_entity_uid for every 'merged' decision on this run."""
        with self.pg.cursor() as cur:
            cur.execute(
                """
                SELECT candidate_record_json, matched_against_entity_uid
                FROM ReviewQueueItem
                WHERE source_run_id = %s AND status = 'merged'
                """,
                (self.params.run_id,),
            )
            rows = cur.fetchall()
        remap: dict[str, str] = {}
        for row in rows:
            candidate = row["candidate_record_json"]
            if isinstance(candidate, str):
                candidate = json.loads(candidate or "{}")
            provisional_uid = candidate.get("entity_uid")
            if provisional_uid:
                remap[provisional_uid] = str(row["matched_against_entity_uid"])
        return remap

    # -- SQL load -------------------------------------------------------

    def _next_int_id(self, table: str, pk_column: str) -> int:
        with self.pg.cursor() as cur:
            cur.execute(f"SELECT COALESCE(MAX({pk_column}), 0) + 1 AS next_id FROM {table}")
            return int(cur.fetchone()["next_id"])

    def _ensure_case_master(self, fields: dict[str, Any]) -> None:
        # process_batch may have inserted a stub CaseMaster to satisfy PipelineRun FK;
        # enrich it from FIR when those fields arrive.
        crime_no = fields.get("crime_no") or f"1{self.params.case_id:017d}"[:18]
        case_no = fields.get("case_no") or str(self.params.case_id)
        with self.pg.cursor() as cur:
            cur.execute("SELECT 1 FROM CaseMaster WHERE CaseMasterID = %s", (self.params.case_id,))
            if cur.fetchone():
                cur.execute(
                    """
                    UPDATE CaseMaster SET
                        CrimeNo = COALESCE(%s, CrimeNo),
                        CaseNo = COALESCE(%s, CaseNo),
                        CrimeRegisteredDate = COALESCE(%s, CrimeRegisteredDate),
                        BriefFacts = COALESCE(%s, BriefFacts)
                    WHERE CaseMasterID = %s
                    """,
                    (
                        fields.get("crime_no"),
                        fields.get("case_no"),
                        fields.get("crime_registered_date"),
                        fields.get("brief_facts"),
                        self.params.case_id,
                    ),
                )
            else:
                # Fallback only when no stub exists (non-demo path). Demo stubs are
                # inserted by process_batch with per-scenario FIR-aligned defaults.
                cur.execute(
                    """
                    INSERT INTO CaseMaster
                        (CaseMasterID, CrimeNo, CaseNo, CrimeRegisteredDate, PolicePersonID, PoliceStationID,
                         CaseCategoryID, GravityOffenceID, CrimeMajorHeadID, CrimeMinorHeadID, CaseStatusID, CourtID, BriefFacts)
                    VALUES (%s, %s, %s, COALESCE(%s, NOW()), 5001, 1001, 1, 3, 101, 1011, 1, 7001, %s)
                    """,
                    (
                        self.params.case_id,
                        crime_no,
                        case_no,
                        fields.get("crime_registered_date"),
                        fields.get("brief_facts"),
                    ),
                )
        self.pg.commit()

    def _ensure_entity(self, table_name: str, pk_value: int, entity_type: str, pole_subtype: str, entity_uid: str | None = None) -> str:
        with self.pg.cursor() as cur:
            cur.execute(
                "SELECT entity_uid FROM EntityMap WHERE sql_table = %s AND sql_pk = %s AND status = 'active' LIMIT 1",
                (table_name, str(pk_value)),
            )
            row = cur.fetchone()
            if row:
                return str(row["entity_uid"])
            uid = entity_uid or str(uuid.uuid4())
            cur.execute(
                """
                INSERT INTO EntityMap(entity_uid, entity_type, pole_subtype, sql_table, sql_pk, status, created_at, updated_at)
                VALUES (%s, %s, %s, %s, %s, 'active', NOW(), NOW())
                """,
                (uid, entity_type, pole_subtype, table_name, str(pk_value)),
            )
        self.pg.commit()
        return uid

    def _write_object_row(self, table: str, group_fields_row: dict[str, Any], evidence_id: int | None) -> dict[str, Any] | None:
        spec = OBJECT_TABLE[table]
        fields = group_fields_row["fields"]
        columns: dict[str, str] = group_fields_row.get("columns") or {}
        identifier_field = group_fields_row.get("identifier_field")
        identifier_value = fields.get(identifier_field) if identifier_field else None
        if not identifier_value:
            return None
        normalized = _norm(spec["identifier_type"], str(identifier_value))

        with self.pg.cursor() as cur:
            cur.execute(f"SELECT {spec['pk']} AS id FROM {table} WHERE {spec['norm']} = %s LIMIT 1", (normalized,))
            existing = cur.fetchone()
            created = existing is None
            if existing:
                object_id = int(existing["id"])
            else:
                cur.execute(
                    f"INSERT INTO {table} ({spec['raw']}, {spec['norm']}, created_at) VALUES (%s, %s, NOW()) RETURNING {spec['pk']} AS id",
                    (str(identifier_value), normalized),
                )
                object_id = int(cur.fetchone()["id"])
        self.pg.commit()

        # Update the remaining descriptive columns of the same group (e.g. holder_name_raw,
        # ifsc, bank_name...), read straight off SchemaField.target_column - no per-table
        # hardcoding, so a group like UPIPayee can share the UPIHandle table with UPIPayer
        # under a different field_name and still land on the right column.
        holder_name_display = None
        other_columns: dict[str, Any] = {}
        for field_name, value in fields.items():
            if field_name == identifier_field or value is None:
                continue
            column = columns.get(field_name)
            if column and column != spec["raw"]:
                other_columns[column] = value
        if other_columns:
            set_clause = ", ".join(f"{col} = %s" for col in other_columns)
            with self.pg.cursor() as cur:
                cur.execute(f"UPDATE {table} SET {set_clause} WHERE {spec['pk']} = %s", (*other_columns.values(), object_id))
            self.pg.commit()
            holder_name_display = other_columns.get("holder_name_raw")

        if evidence_id:
            # source_evidence_id is only stamped on rows THIS pipeline created -- see
            # _stamp_source_evidence for why touching a pre-existing row is destructive.
            sets: list[str] = []
            values: list[Any] = []
            if created:
                sets.append("source_evidence_id = COALESCE(source_evidence_id, %s)")
                values.append(evidence_id)
            if table == "Account":
                sets.append("linked_case_id = COALESCE(linked_case_id, %s)")
                values.append(self.params.case_id)
            if sets:
                with self.pg.cursor() as cur:
                    cur.execute(
                        f"UPDATE {table} SET {', '.join(sets)} WHERE {spec['pk']} = %s",
                        (*values, object_id),
                    )
                self.pg.commit()

        entity_uid = self._ensure_entity(table, object_id, "Object", table)
        display_name = f"{holder_name_display} · {identifier_value}" if holder_name_display else str(identifier_value)
        return {
            "entity_uid": entity_uid, "table": table, "pk": object_id,
            "raw": {"holder_name_display": holder_name_display, "display_name": display_name, **fields},
        }

    def _write_person_row(
        self, table: str, row: dict[str, Any], idx: int, uid_remap: dict[str, str], provisional_uids: dict[str, str], provisional_key: str
    ) -> dict[str, Any]:
        spec = PERSON_TABLE[table]
        fields = row["fields"]
        columns: dict[str, str] = row.get("columns") or {}

        provisional_uid = provisional_uids.get(f"{provisional_key}#{idx}")
        final_uid = uid_remap.get(provisional_uid, provisional_uid) if provisional_uid else None

        # Person tables have no natural key to dedup on (unlike Account/UPI/Phone/
        # Device, which look up by normalized value in _insert_or_get_identifier
        # before inserting) -- entity_uid is the only stable handle. If this
        # mention already resolved to an existing entity (reviewed match, or a
        # retry of an already-loaded run), reuse that row instead of minting a
        # second SQL row for the same person.
        if final_uid:
            with self.pg.cursor() as cur:
                cur.execute(
                    "SELECT sql_table, sql_pk FROM EntityMap WHERE entity_uid = %s AND status = 'active'",
                    (final_uid,),
                )
                existing = cur.fetchone()
            if existing and existing["sql_table"] == table:
                pk_value = int(existing["sql_pk"])
                return {"entity_uid": final_uid, "table": table, "pk": pk_value, "raw": {"display_name": _display_name(fields), **fields}}

        pk_value = self._next_int_id(table, spec["pk"])
        insert_items: list[tuple[str, Any]] = []
        for name, value in fields.items():
            if value is None or name not in columns:
                continue
            column = columns[name]
            if column in self._LOOKUP_TABLES:
                resolved = self._resolve_lookup_id(column, str(value))
                if resolved is None:
                    continue
                insert_items.append((column, resolved))
            else:
                insert_items.append((column, value))

        # An unknown accused (money-laundering/task-scam FIRs often name no offender) leaves
        # the NOT NULL name column absent. Historical rows store "Unknown"; do the same rather
        # than crash the whole Phase B load on a NotNullViolation.
        name_col = spec["name_column"]
        if not any(c == name_col for c, _ in insert_items):
            insert_items.append((name_col, "Unknown"))

        with self.pg.cursor() as cur:
            col_sql = ", ".join([spec["pk"], "CaseMasterID"] + [c for c, _ in insert_items])
            placeholders = ", ".join(["%s"] * (2 + len(insert_items)))
            cur.execute(
                f"INSERT INTO {table} ({col_sql}) VALUES ({placeholders})",
                (pk_value, self.params.case_id, *[v for _, v in insert_items]),
            )
        self.pg.commit()

        entity_uid = self._ensure_entity(table, pk_value, "Person", table, entity_uid=final_uid)

        return {"entity_uid": entity_uid, "table": table, "pk": pk_value, "raw": {"display_name": _display_name(fields), **fields}}

    def _write_evidence_row(self, file_entry: dict[str, Any]) -> tuple[int, dict[str, Any]]:
        with self.pg.cursor() as cur:
            cur.execute(
                """
                INSERT INTO Evidence(case_id, doc_type, file_ref, original_filename, extraction_status, schema_id_used, uploaded_by, upload_ts, created_at)
                VALUES (%s, %s, %s, %s, 'success', %s, 'catalyst-function', NOW(), NOW())
                RETURNING evidence_id
                """,
                (self.params.case_id, file_entry["doc_type"], file_entry["key"], file_entry["filename"], file_entry["schema_id"]),
            )
            evidence_id = int(cur.fetchone()["evidence_id"])
        self.pg.commit()
        entity_uid = self._ensure_entity("Evidence", evidence_id, "Object", "Evidence")
        return evidence_id, {
            "entity_uid": entity_uid, "table": "Evidence", "pk": evidence_id,
            "raw": {"display_name": file_entry["filename"], "doc_type": file_entry["doc_type"]},
        }

    def _write_investigation_report_row(self, fields: dict[str, Any], schema_id: int) -> dict[str, Any]:
        with self.pg.cursor() as cur:
            cur.execute(
                """
                INSERT INTO InvestigationReport(case_id, report_date, findings_narrative, status, schema_id_used, created_at)
                VALUES (%s, COALESCE(%s, CURRENT_DATE), %s, 'draft', %s, NOW())
                RETURNING report_id
                """,
                (self.params.case_id, fields.get("report_date"), fields.get("findings_narrative"), schema_id),
            )
            report_id = int(cur.fetchone()["report_id"])
        self.pg.commit()
        entity_uid = self._ensure_entity("InvestigationReport", report_id, "Event", "InvestigationReport")
        report_date = fields.get("report_date")
        display_name = f"Investigation Report ({report_date})" if report_date else "Investigation Report"
        return {
            "entity_uid": entity_uid, "table": "InvestigationReport", "pk": report_id,
            "raw": {**fields, "display_name": display_name},
        }

    def _resolve_act_section(self, act_raw: str, section_raw: str) -> tuple[str, str] | None:
        """The LLM returns free text like "IT Act" / "S.66C"; ActSectionAssociation's
        FK needs the master (ActCode, SectionCode) pair, e.g. ("ITACT", "66C")."""
        cache = self._lookup_cache.setdefault("Act", {})
        if not cache:
            with self.pg.cursor() as cur:
                cur.execute("SELECT ActCode AS code, ActDescription AS descr, ShortName AS short FROM Act")
                for row in cur.fetchall():
                    for key in (row["code"], row["short"], row["descr"]):
                        if key:
                            cache[str(key).strip().lower()] = row["code"]

        act_key = act_raw.strip().lower()
        act_code = cache.get(act_key)
        if not act_code:
            best_name, best_score = None, 0.0
            for name in cache:
                score = fuzz.ratio(act_key, name)
                if score > best_score:
                    best_name, best_score = name, score
            if best_name and best_score >= 80:
                act_code = cache[best_name]
        if not act_code:
            logger.warning("act_section: no Act match for %r — skipping", act_raw)
            return None

        section_code = re.sub(r"(?i)^\s*(sec(tion)?\.?|s\.)\s*", "", section_raw).strip()
        with self.pg.cursor() as cur:
            cur.execute("SELECT 1 FROM Section WHERE ActCode = %s AND SectionCode = %s", (act_code, section_code))
            if cur.fetchone():
                return act_code, section_code
        logger.warning("act_section: no Section match for act=%s section=%r — skipping", act_code, section_raw)
        return None

    def _write_act_section_rows(self, rows: list[dict[str, Any]]) -> None:
        for row in rows:
            fields = row["fields"]
            act_raw, section_raw = fields.get("act_code"), fields.get("section_code")
            if not act_raw or not section_raw:
                continue
            resolved = self._resolve_act_section(str(act_raw), str(section_raw))
            if resolved is None:
                continue
            act_code, section_code = resolved
            with self.pg.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO ActSectionAssociation(CaseMasterID, ActCode, SectionCode)
                    VALUES (%s, %s, %s)
                    ON CONFLICT DO NOTHING
                    """,
                    (self.params.case_id, act_code, section_code),
                )
            self.pg.commit()

    def _write_transaction_rows(self, rows: list[dict[str, Any]], produced: dict[str, list[dict[str, Any]]], evidence_id: int | None) -> list[dict[str, Any]]:
        self_account = produced.get("BankStatement", [None])[0] if produced.get("BankStatement") else None
        payer = produced.get("UPIPayer", [None])[0] if produced.get("UPIPayer") else None
        payee = produced.get("UPIPayee", [None])[0] if produced.get("UPIPayee") else None

        edges: list[dict[str, Any]] = []
        for row in rows:
            fields = row["fields"]
            amount = fields.get("amount")
            if amount is None:
                continue
            txn_ts = _normalize_timestamp(fields.get("txn_timestamp"))
            mode, utr_ref, direction = fields.get("mode"), fields.get("utr_ref"), str(fields.get("direction") or "").lower()

            from_account_id = to_account_id = from_upi_id = to_upi_id = None
            from_uid = to_uid = None

            if self_account and ("counterparty_account" in fields or "counterparty_upi" in fields):
                counterparty = None
                if fields.get("counterparty_account"):
                    cp_id, cp_uid = self._insert_or_get_identifier("Account", str(fields["counterparty_account"]), evidence_id)
                    counterparty = {"pk": cp_id, "uid": cp_uid, "table": "account"}
                elif fields.get("counterparty_upi"):
                    cp_id, cp_uid = self._insert_or_get_identifier("UPIHandle", str(fields["counterparty_upi"]), evidence_id)
                    counterparty = {"pk": cp_id, "uid": cp_uid, "table": "upi"}

                self_side = {"pk": self_account["pk"], "uid": self_account["entity_uid"], "table": "account"}
                if direction in ("debit", "dr", "out"):
                    from_side, to_side = self_side, counterparty
                else:
                    from_side, to_side = counterparty, self_side

                if from_side:
                    from_uid = from_side["uid"]
                    if from_side["table"] == "account":
                        from_account_id = from_side["pk"]
                    else:
                        from_upi_id = from_side["pk"]
                if to_side:
                    to_uid = to_side["uid"]
                    if to_side["table"] == "account":
                        to_account_id = to_side["pk"]
                    else:
                        to_upi_id = to_side["pk"]
            elif payer and payee:
                from_upi_id, to_upi_id = payer["pk"], payee["pk"]
                from_uid, to_uid = payer["entity_uid"], payee["entity_uid"]
            else:
                continue

            with self.pg.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO Transaction
                        (from_account_id, from_upi_id, to_account_id, to_upi_id, amount, txn_timestamp, mode, utr_ref, direction, source_evidence_id, created_at)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s, NOW())
                    """,
                    (from_account_id, from_upi_id, to_account_id, to_upi_id, amount, txn_ts, mode, utr_ref, direction or None, evidence_id),
                )
            self.pg.commit()

            if from_uid and to_uid:
                edges.append({"from_uid": from_uid, "to_uid": to_uid, "amount": amount, "txn_timestamp": str(txn_ts), "utr_ref": utr_ref})
        return edges

    def _stamp_source_evidence(self, table: str, spec: dict[str, Any], object_id: int, evidence_id: int) -> None:
        """Record which demo evidence file first produced this row.

        Only ever called for rows this pipeline just created. A pre-existing row may be
        a *historical* object that a live upload legitimately references -- the planted
        aggregation account is exactly that, and it is the point of the demo. Historical
        rows carry source_evidence_id = NULL, so COALESCE(source_evidence_id, <demo>)
        would happily claim them for the demo; reset_demo_data.py then deletes every
        object whose source_evidence_id belongs to a demo case, taking the planted
        account with it and quietly breaking the demo's own links on re-run.
        """
        with self.pg.cursor() as cur:
            cur.execute(
                f"UPDATE {table} SET source_evidence_id = COALESCE(source_evidence_id, %s) WHERE {spec['pk']} = %s",
                (evidence_id, object_id),
            )
        self.pg.commit()

    def _insert_or_get_identifier(self, table: str, raw_value: str, evidence_id: int | None) -> tuple[int, str]:
        spec = OBJECT_TABLE[table]
        normalized = _norm(spec["identifier_type"], raw_value)
        with self.pg.cursor() as cur:
            cur.execute(f"SELECT {spec['pk']} AS id FROM {table} WHERE {spec['norm']} = %s LIMIT 1", (normalized,))
            existing = cur.fetchone()
            created = existing is None
            if existing:
                object_id = int(existing["id"])
            else:
                cur.execute(
                    f"INSERT INTO {table} ({spec['raw']}, {spec['norm']}, created_at) VALUES (%s, %s, NOW()) RETURNING {spec['pk']} AS id",
                    (raw_value, normalized),
                )
                object_id = int(cur.fetchone()["id"])
        self.pg.commit()
        if evidence_id and created:
            self._stamp_source_evidence(table, spec, object_id, evidence_id)
        entity_uid = self._ensure_entity(table, object_id, "Object", table)
        return object_id, entity_uid

    def _load_sql_groups(self, file_entry: dict[str, Any], uid_remap: dict[str, str]) -> tuple[dict[str, list[dict[str, Any]]], int | None]:
        doc_type = file_entry["doc_type"]
        groups_raw = file_entry["groups"]
        provisional_uids = file_entry.get("provisional_uids", {})
        produced: dict[str, list[dict[str, Any]]] = {}
        txn_edges: list[dict[str, Any]] = []

        if doc_type == "FIR":
            case_fields = (groups_raw.get("CaseMaster") or [{"fields": {}}])[0]["fields"]
            self._ensure_case_master(case_fields)

        case_entity_uid = self._ensure_entity(
            "CaseMaster", self.params.case_id, "Event", "Case",
        )
        with self.pg.cursor() as cur:
            cur.execute("SELECT CrimeNo FROM CaseMaster WHERE CaseMasterID = %s", (self.params.case_id,))
            crime_no_row = cur.fetchone()
        case_display_name = (crime_no_row or {}).get("crimeno") or f"Case {self.params.case_id}"
        produced["CaseMaster"] = [{
            "entity_uid": case_entity_uid, "table": "CaseMaster", "pk": self.params.case_id,
            "raw": {"display_name": case_display_name},
        }]

        evidence_id: int | None = None
        if doc_type.startswith("EVIDENCE"):
            evidence_id, evidence_node = self._write_evidence_row(file_entry)
            produced["Evidence"] = [evidence_node]

        if doc_type == "IR":
            ir_fields = (groups_raw.get("InvestigationReport") or [{"fields": {}}])[0]["fields"]
            produced["InvestigationReport"] = [self._write_investigation_report_row(ir_fields, file_entry["schema_id"])]

        for group_name, rows in groups_raw.items():
            if group_name in ("CaseMaster", "InvestigationReport", "Transaction", "ActSection", "Evidence"):
                continue
            if not rows:
                continue
            # Dispatch on the group's *target_table*, not its group_name — several
            # distinct groups (MentionedAccount, BankStatement, ...) all target the
            # same physical Account/UPIHandle/PhoneNumber/Device table.
            target_table = rows[0].get("target_table")
            if target_table in OBJECT_TABLE:
                produced_rows = []
                for row in rows:
                    written = self._write_object_row(target_table, row, evidence_id)
                    if written:
                        produced_rows.append(written)
                produced[group_name] = produced_rows
            elif target_table in PERSON_TABLE:
                produced[group_name] = [
                    self._write_person_row(target_table, row, idx, uid_remap, provisional_uids, provisional_key=group_name)
                    for idx, row in enumerate(rows)
                    if any(v for v in row["fields"].values())
                ]

        if "ActSection" in groups_raw:
            self._write_act_section_rows(groups_raw["ActSection"])

        if "Transaction" in groups_raw:
            txn_edges = self._write_transaction_rows(groups_raw["Transaction"], produced, evidence_id)
        produced["__txn_edges__"] = txn_edges  # type: ignore[assignment]

        return produced, evidence_id

    # -- Graph load -------------------------------------------------------

    def _load_graph(self, schema_id: int, produced: dict[str, Any]) -> list[str]:
        neo4j_uri = _env("NEO4J_URI")
        neo4j_user = _env("NEO4J_USERNAME")
        if not neo4j_uri or not _env("NEO4J_PASSWORD"):
            logger.warning("load_graph: NEO4J_URI or NEO4J_PASSWORD not set — skipping graph load")
            return []
        try:
            from neo4j import GraphDatabase
        except Exception as exc:
            logger.error("load_graph: neo4j package not importable: %s", exc)
            return []

        txn_edges = produced.pop("__txn_edges__", [])
        edges = self._build_edges(schema_id, produced, txn_edges)
        logger.debug("load_graph: uri=%s nodes=%d edges=%d", neo4j_uri,
                     sum(len(v) for v in produced.values()), len(edges))

        touched: list[str] = []
        try:
            driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, _env("NEO4J_PASSWORD")))
        except Exception as exc:
            logger.error("load_graph: driver creation failed uri=%s: %s", neo4j_uri, exc)
            return []
        try:
            with driver.session() as session:
                for group_name, rows in produced.items():
                    for node in rows:
                        uid = node.get("entity_uid")
                        if not uid:
                            continue
                        label = _sanitize_label(group_name)
                        display_name = (node.get("raw") or {}).get("display_name")
                        # MERGE must key on entity_uid alone, not entity_uid+label --
                        # Cypher's MERGE (n:Label {prop: val}) only reuses a node if
                        # BOTH the label and property already match. The same real
                        # account gets mentioned under different schema group names
                        # (e.g. "MentionedAccount" from a FIR, "BankStatement" from a
                        # bank-statement evidence file) -- with a label in the MERGE
                        # pattern, the second write couldn't find the first node and
                        # silently created a duplicate with the same entity_uid.
                        # Confirmed live: a single account ended up as two disconnected
                        # Neo4j nodes. Label-less MERGE + SET n:Label adds the label to
                        # whichever node already owns that entity_uid, idempotently.
                        # origin is first-writer-wins, NOT unconditional. The planted
                        # shared objects (aggregation account, controller UPI/IMEI) are
                        # historical nodes that a live upload legitimately touches --
                        # that shared node IS the demo's payoff. Overwriting origin to
                        # 'demo' here made reset_demo_data.py's
                        # `MATCH (n {origin:'demo'}) DETACH DELETE n` delete the planted
                        # historical node and every historical edge hanging off it, so
                        # the demo silently stopped finding its own links on the second
                        # run. coalesce keeps 'historical' on shared nodes; genuinely
                        # new live-only nodes still get 'demo' and stay wipeable.
                        session.run(
                            f"""
                            MERGE (n {{entity_uid: $uid}})
                            SET n:{label},
                                n.origin = coalesce(n.origin, 'demo'), n.run_id = $run_id, n.case_id = $case_id,
                                n.updated_at = datetime($ts),
                                n.display_name = coalesce($display_name, n.display_name)
                            """,
                            uid=uid, run_id=self.params.run_id, case_id=self.params.case_id, ts=_utc_now(),
                            display_name=display_name,
                        )
                        touched.append(uid)

                for edge in edges:
                    props = {k: v for k, v in (edge.get("props") or {}).items() if v is not None}
                    if edge["type"] == "TRANSACTED_WITH":
                        txn_key = hashlib.sha1(
                            f"{edge['from_uid']}|{edge['to_uid']}|{props.get('amount')}|{props.get('txn_timestamp')}".encode("utf-8")
                        ).hexdigest()
                        session.run(
                            """
                            MATCH (a {entity_uid: $from_uid}), (b {entity_uid: $to_uid})
                            MERGE (a)-[r:TRANSACTED_WITH {txn_key: $txn_key}]->(b)
                            SET r += $props, r.origin = 'demo', r.run_id = $run_id
                            """,
                            from_uid=edge["from_uid"], to_uid=edge["to_uid"], txn_key=txn_key, props=props, run_id=self.params.run_id,
                        )
                    else:
                        rel_type = _sanitize_label(edge["type"])
                        session.run(
                            f"""
                            MATCH (a {{entity_uid: $from_uid}}), (b {{entity_uid: $to_uid}})
                            MERGE (a)-[r:{rel_type}]->(b)
                            SET r += $props, r.origin = 'demo', r.run_id = $run_id
                            """,
                            from_uid=edge["from_uid"], to_uid=edge["to_uid"], props=props, run_id=self.params.run_id,
                        )
            logger.info("load_graph: ok nodes_touched=%d edges=%d", len(touched), len(edges))
        except Exception as exc:
            logger.error("load_graph: neo4j session error uri=%s: %s", neo4j_uri, exc)
            raise
        finally:
            driver.close()
        return touched

    def _build_edges(self, schema_id: int, produced: dict[str, list[dict[str, Any]]], txn_edges: list[dict[str, Any]]) -> list[dict[str, Any]]:
        with self.pg.cursor() as cur:
            cur.execute(
                "SELECT from_group, to_group, relationship_type, direction, fixed_edge_properties FROM SchemaRelationship WHERE schema_id = %s",
                (schema_id,),
            )
            rels = cur.fetchall()

        edges: list[dict[str, Any]] = []
        for rel in rels:
            # TRANSACTED_WITH is always sourced from txn_edges (real amount/timestamp/utr_ref
            # from the actual Transaction rows), never from a generic node cartesian product -
            # that would create spurious props-less duplicate edges alongside the real ones.
            if rel["relationship_type"] == "TRANSACTED_WITH":
                continue

            from_nodes = [n for n in produced.get(rel["from_group"], []) if n.get("entity_uid")]
            to_nodes = [n for n in produced.get(rel["to_group"], []) if n.get("entity_uid")]
            if not from_nodes or not to_nodes:
                continue
            fixed_props = rel["fixed_edge_properties"] or {}
            if isinstance(fixed_props, str):
                fixed_props = json.loads(fixed_props or "{}")

            if rel["relationship_type"] == "OWNS":
                edges.extend(self._owns_edges(from_nodes, to_nodes, fixed_props))
                continue

            for from_node in from_nodes:
                for to_node in to_nodes:
                    edges.append({"from_uid": from_node["entity_uid"], "to_uid": to_node["entity_uid"], "type": rel["relationship_type"], "props": dict(fixed_props)})

        for txn in txn_edges:
            edges.append({
                "from_uid": txn["from_uid"], "to_uid": txn["to_uid"], "type": "TRANSACTED_WITH",
                "props": {"amount": txn["amount"], "txn_timestamp": txn["txn_timestamp"], "utr_ref": txn.get("utr_ref")},
            })
        return edges

    def _owns_edges(self, person_nodes: list[dict[str, Any]], object_nodes: list[dict[str, Any]], fixed_props: dict[str, Any]) -> list[dict[str, Any]]:
        # ponytail: holder-name fuzzy match when available, else only link when the
        # from_group has exactly one candidate (unambiguous). Otherwise skip rather
        # than guess. Ceiling: no cross-document holder linkage. Upgrade path: run
        # this pass again in Phase B against the full case's persons, not just this file's.
        edges = []
        for obj in object_nodes:
            holder_name = str((obj.get("raw") or {}).get("holder_name_display") or "").strip()
            best, best_score = None, 0.0
            if holder_name:
                for person in person_nodes:
                    pname = str((person.get("raw") or {}).get("display_name") or "").strip()
                    if not pname:
                        continue
                    score = fuzz.token_set_ratio(holder_name, pname) / 100.0
                    if score > best_score:
                        best_score, best = score, person
            if best and best_score >= 0.72:
                edges.append({"from_uid": best["entity_uid"], "to_uid": obj["entity_uid"], "type": "OWNS", "props": dict(fixed_props)})
            elif len(person_nodes) == 1:
                edges.append({"from_uid": person_nodes[0]["entity_uid"], "to_uid": obj["entity_uid"], "type": "OWNS", "props": dict(fixed_props)})
        return edges

    # -- Vector load -------------------------------------------------------

    def _load_vector(self, file_entry: dict[str, Any], entity_uids: list[str]) -> list[str]:
        raw_text = file_entry.get("raw_text") or ""
        if not raw_text.strip():
            logger.warning("load_vector: skipping file=%s — raw_text is empty", file_entry.get("filename"))
            return []
        pinecone_key = _env("PINECONE_API_KEY")
        if not pinecone_key:
            logger.warning("load_vector: PINECONE_API_KEY not set — skipping vector load")
            return []
        try:
            import boto3
            from pinecone import Pinecone
        except Exception as exc:
            logger.error("load_vector: missing dependency: %s", exc)
            return []

        region = _env("AWS_DEFAULT_REGION") or _env("AWS_REGION") or "us-east-1"
        embed_model = _env("BEDROCK_EMBEDDING_MODEL", "amazon.titan-embed-text-v1")
        logger.debug("load_vector: bedrock region=%s model=%s text_len=%d", region, embed_model, len(raw_text))
        try:
            bedrock = boto3.client("bedrock-runtime", region_name=region)
            embed_resp = bedrock.invoke_model(
                modelId=embed_model,
                body=json.dumps({"inputText": raw_text[:40000]}),
            )
            values = json.loads(embed_resp["body"].read().decode("utf-8")).get("embedding")
        except Exception as exc:
            logger.error("load_vector: bedrock invoke_model failed region=%s model=%s: %s", region, embed_model, exc)
            return []
        if not values:
            logger.error("load_vector: bedrock returned no embedding for file=%s", file_entry.get("filename"))
            return []

        chunk_id = f"demo::{self.params.run_id}::{file_entry['filename']}::0"
        metadata = {
            "case_id": self.params.case_id,
            "doc_type": file_entry["doc_type"],
            "run_id": self.params.run_id,
            "origin": "demo",
            "graph_node_ids": entity_uids,
            "source": file_entry["key"],
            "chunk_index": 0,
        }
        index_name = _env("PINECONE_INDEX", "ksp-crime-intel")
        logger.debug("load_vector: upserting chunk_id=%s to index=%s", chunk_id, index_name)
        try:
            pc = Pinecone(api_key=pinecone_key)
            index = pc.Index(index_name)
            index.upsert(vectors=[{"id": chunk_id, "values": values, "metadata": metadata}])
            logger.info("load_vector: ok chunk_id=%s index=%s", chunk_id, index_name)
        except Exception as exc:
            logger.error("load_vector: pinecone upsert failed index=%s chunk_id=%s: %s", index_name, chunk_id, exc)
            return []
        return [chunk_id]

    def _writeback_links(self, chunk_ids: list[str], entity_uids: list[str]) -> None:
        if not (chunk_ids and entity_uids and _env("NEO4J_URI") and _env("NEO4J_PASSWORD")):
            logger.debug("writeback_links: skipping — no chunk_ids or no neo4j config")
            return
        try:
            from neo4j import GraphDatabase

            driver = GraphDatabase.driver(_env("NEO4J_URI"), auth=(_env("NEO4J_USERNAME"), _env("NEO4J_PASSWORD")))
            try:
                with driver.session() as session:
                    for uid in entity_uids:
                        session.run(
                            """
                            MATCH (n {entity_uid: $uid})
                            SET n.source_chunks = coalesce(n.source_chunks, []) + $chunk_ids
                            """,
                            uid=uid, chunk_ids=chunk_ids,
                        )
                logger.debug("writeback_links: ok uids=%d chunks=%d", len(entity_uids), len(chunk_ids))
            finally:
                driver.close()
        except Exception as exc:
            logger.warning("writeback_links: failed (non-fatal) — %s", exc)


def _display_name(fields: dict[str, Any]) -> str | None:
    """Generic person-name lookup: the first non-empty field whose name looks like a name.
    Works for accused_name/victim_name/complainant_name/etc. without hardcoding group names."""
    for key, value in fields.items():
        if value and "name" in key.lower():
            return str(value)
    return None
